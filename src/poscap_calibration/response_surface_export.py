"""応答曲面用データを解析Excelの「csv出力用」シートからCSVへ出力する。

せん断試験(手順書1.2.8)と安息角試験(手順書2.2.7)で共通の処理。
条件数が9から27などへ増えても対応できるよう、行数は固定せずシートの
使用範囲をそのまま書き出す。
"""

from __future__ import annotations

import csv
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

#: 応答曲面用データが並ぶシート名(せん断試験・安息角試験で共通)。
RESPONSE_SURFACE_SHEET_NAME = "csv出力用"


class ResponseSurfaceExportError(RuntimeError):
    """応答曲面用データの出力に失敗した場合に送出する例外。"""


def is_recalculation_available() -> bool:
    """Excel COMによる再計算が利用できるかを返す。"""
    try:
        import win32com.client  # noqa: F401
    except ImportError:
        return False
    return True


def export_response_surface_data(
    workbook_path: str | Path,
    output_csv_path: str | Path,
    *,
    recalculate: bool = True,
    allow_incomplete: bool = False,
    sheet_name: str = RESPONSE_SURFACE_SHEET_NAME,
) -> dict[str, Any]:
    """「csv出力用」シート全体を応答曲面用CSVとして出力する。

    数式エラーが残っている場合は既定で失敗させ、未完了のデータを応答曲面へ渡さない。
    """
    path = Path(workbook_path).resolve()
    if not path.is_file():
        raise FileNotFoundError(f"解析Excelが見つかりません: {path}")

    result: dict[str, Any] = {
        "workbook": str(path),
        "sheet": sheet_name,
        "recalculated": False,
    }

    if recalculate:
        result["recalculated"] = _recalculate_and_save(path)
        if not result["recalculated"]:
            result["recalculation_required"] = True
            result["note"] = (
                "Excel COM(pywin32)が利用できないため再計算していません。"
                "出力値は保存済みのキャッシュ値です。"
            )

    rows = _read_used_range(path, sheet_name)
    if not rows:
        raise ResponseSurfaceExportError(f"シート '{sheet_name}' にデータがありません: {path}")

    error_cells = [
        f"R{row_index}C{column_index}={value}"
        for row_index, row in enumerate(rows, start=1)
        for column_index, value in enumerate(row, start=1)
        if isinstance(value, str) and value.startswith("#")
    ]
    if error_cells and not allow_incomplete:
        raise ResponseSurfaceExportError(
            f"シート '{sheet_name}' に数式エラーが残っています: "
            f"{error_cells[:10]}。全条件の解析完了後に実行してください。"
        )

    output_path = Path(output_csv_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    # 日本語見出しをExcelでそのまま開けるようBOM付きUTF-8で出力する。
    with output_path.open("w", encoding="utf-8-sig", newline="") as csv_file:
        writer = csv.writer(csv_file)
        for row in rows:
            writer.writerow(["" if value is None else value for value in row])

    result.update(
        {
            "output_csv": str(output_path.resolve()),
            "row_count": len(rows),
            "column_count": len(rows[0]),
            "data_row_count": max(len(rows) - 1, 0),
        }
    )
    if error_cells:
        result["formula_errors"] = error_cells
    return result


def _read_used_range(workbook_path: Path, sheet_name: str) -> list[list[Any]]:
    """シートの使用範囲を読み取り、末尾の空行・空列を取り除いて返す。"""
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        if sheet_name not in workbook.sheetnames:
            raise ResponseSurfaceExportError(
                f"シート '{sheet_name}' がありません: {workbook_path}。"
                f"利用可能なシート: {workbook.sheetnames}"
            )
        worksheet = workbook[sheet_name]
        rows = [list(row) for row in worksheet.iter_rows(values_only=True)]
    finally:
        workbook.close()

    while rows and all(value is None for value in rows[-1]):
        rows.pop()
    if not rows:
        return []

    last_column = 0
    for row in rows:
        for column_index, value in enumerate(row, start=1):
            if value is not None:
                last_column = max(last_column, column_index)
    return [row[:last_column] for row in rows]


def _recalculate_and_save(workbook_path: Path) -> bool:
    """ブックを開いて完全再計算し保存する。COMが使えない場合はFalseを返す。"""
    if not is_recalculation_available():
        return False

    import pythoncom
    import win32com.client

    pythoncom.CoInitialize()
    excel = None
    workbook = None
    try:
        excel = win32com.client.DispatchEx("Excel.Application")
        excel.Visible = False
        excel.DisplayAlerts = False
        workbook = excel.Workbooks.Open(str(workbook_path))
        excel.Application.CalculateFullRebuild()
        workbook.Save()
    finally:
        if workbook is not None:
            workbook.Close(SaveChanges=False)
        if excel is not None:
            excel.Quit()
        pythoncom.CoUninitialize()
    return True
