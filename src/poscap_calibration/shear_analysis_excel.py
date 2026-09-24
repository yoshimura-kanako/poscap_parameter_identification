"""解析Excel(shear_analysis.xlsx)への本せん断時系列の書込みと計算結果の取得。

グラフ・書式・配列数式を確実に保持するため、書込みと再計算はExcel COM経由で行う。
COMが使えない環境では書込みを行わず、再計算が必要である旨を呼び出し側へ返す。
"""

from __future__ import annotations

import re
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from poscap_calibration.response_surface_export import (
    RESPONSE_SURFACE_SHEET_NAME,
    ResponseSurfaceExportError,
    export_response_surface_data,
    is_recalculation_available,
)

__all__ = [
    "RESPONSE_SURFACE_SHEET_NAME",
    "ShearAnalysisExcelError",
    "export_response_surface_data",
    "find_time_series_columns",
    "get_correction_factor",
    "get_peak_ratio",
    "get_shear_sheet_names",
    "get_steady_slope_cells",
    "is_recalculation_available",
    "update_and_recalculate",
]

SETTINGS_SHEET_NAME = "設定・集計"
_CONDITION_HEADER = "条件"
#: 設定・集計シートで荷重[kPa]と本せん断シート名列を対応付ける見出し。
_SHEAR_SHEET_HEADERS: dict[float, str] = {
    3.0: "本せん断3kPa",
    5.0: "本せん断5kPa",
    7.0: "本せん断7kPa",
}
_STEADY_SLOPE_HEADER = "定常傾き"
_STEADY_SLOPE_RAW_HEADER = "補正前"
_STEADY_SLOPE_CORRECTED_HEADER = "補正後"
_CORRECTION_FACTOR_LABEL = "補正係数"
_PEAK_RATIO_LABEL = "ピーク比"

# 本せん断シートのせん断応力列は ``=IF(ISBLANK(F2), NA(), -F2*$B$7/1000)`` の形で
# Moment Y列を参照する。この相対参照からMoment Y列を特定する(シートに見出し行が無いため)。
_SHEAR_STRESS_FORMULA_MARKER = "ISBLANK("
_RELATIVE_REFERENCE_PATTERN = re.compile(r"(?<![$A-Z0-9])([A-Z]{1,3})(\d+)(?![(\d])")

_HEADER_SEARCH_ROWS = 20
_HEADER_SEARCH_COLUMNS = 30
_MAX_CONDITION_ROWS = 200

#: Excel COMが数式エラーを返す際の値。数値として扱うと誤った結果になるため名前へ変換する。
_EXCEL_ERROR_VALUES: dict[int, str] = {
    -2146826288: "#NULL!",
    -2146826281: "#DIV/0!",
    -2146826273: "#VALUE!",
    -2146826265: "#REF!",
    -2146826259: "#NAME?",
    -2146826252: "#NUM!",
    -2146826246: "#N/A",
}


class ShearAnalysisExcelError(ResponseSurfaceExportError):
    """解析Excelの解析・書込みに失敗した場合に送出する例外。"""


def _read_cells(workbook_path: Path, sheet_name: str, max_row: int) -> list[tuple[Any, ...]]:
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        if sheet_name not in workbook.sheetnames:
            raise ShearAnalysisExcelError(
                f"シート '{sheet_name}' がありません: {workbook_path}"
            )
        worksheet = workbook[sheet_name]
        return list(
            worksheet.iter_rows(
                min_row=1, max_row=max_row, max_col=_HEADER_SEARCH_COLUMNS, values_only=True
            )
        )
    finally:
        workbook.close()


def _cell(rows: list[tuple[Any, ...]], row_index: int, column_index: int) -> Any:
    if row_index >= len(rows) or column_index >= len(rows[row_index]):
        return None
    return rows[row_index][column_index]


def _find_header_column(
    rows: list[tuple[Any, ...]], row_index: int, header: str
) -> int | None:
    if row_index >= len(rows):
        return None
    matches = [
        column_index
        for column_index, value in enumerate(rows[row_index])
        if isinstance(value, str) and value.strip() == header
    ]
    if len(matches) > 1:
        raise ShearAnalysisExcelError(
            f"見出し '{header}' が複数見つかりました(列: {[m + 1 for m in matches]})。"
        )
    return matches[0] if matches else None


def _locate_condition_row(rows: list[tuple[Any, ...]], condition_id: int) -> tuple[int, int]:
    """``条件``見出しの列と、指定条件のデータ行(0始まり)を返す。"""
    for row_index in range(min(_HEADER_SEARCH_ROWS, len(rows))):
        for column_index, value in enumerate(rows[row_index]):
            if not isinstance(value, str) or value.strip() != _CONDITION_HEADER:
                continue
            for data_row in range(row_index + 2, min(len(rows), row_index + _MAX_CONDITION_ROWS)):
                cell_value = _cell(rows, data_row, column_index)
                if isinstance(cell_value, (int, float)) and int(cell_value) == condition_id:
                    return column_index, data_row
    raise ShearAnalysisExcelError(
        f"「{SETTINGS_SHEET_NAME}」シートに条件{condition_id}の行が見つかりません。"
    )


def get_shear_sheet_names(workbook_path: str | Path, condition_id: int) -> dict[float, str]:
    """条件番号に対応する本せん断3/5/7 kPaのシート名を「設定・集計」から取得する。"""
    path = Path(workbook_path)
    if not path.is_file():
        raise FileNotFoundError(f"解析Excelが見つかりません: {path}")

    rows = _read_cells(path, SETTINGS_SHEET_NAME, _HEADER_SEARCH_ROWS + _MAX_CONDITION_ROWS)
    _, condition_row = _locate_condition_row(rows, condition_id)

    sheet_names: dict[float, str] = {}
    for load_kpa, header in _SHEAR_SHEET_HEADERS.items():
        column_index = None
        for header_row in range(min(_HEADER_SEARCH_ROWS, len(rows))):
            column_index = _find_header_column(rows, header_row, header)
            if column_index is not None:
                break
        if column_index is None:
            raise ShearAnalysisExcelError(
                f"「{SETTINGS_SHEET_NAME}」シートに見出し '{header}' が見つかりません。"
            )

        value = _cell(rows, condition_row, column_index)
        if value is None:
            raise ShearAnalysisExcelError(
                f"条件{condition_id}の '{header}' シート名が空です。"
                "Excelを開いて保存し、数式の計算結果を確定させてください。"
            )
        sheet_names[load_kpa] = _normalize_sheet_name(value)

    available = set(load_workbook(path, read_only=True).sheetnames)
    missing = {
        load_kpa: name for load_kpa, name in sheet_names.items() if name not in available
    }
    if missing:
        raise ShearAnalysisExcelError(
            f"条件{condition_id}の本せん断シートが存在しません: {missing}。"
            f"利用可能なシート: {sorted(available)}"
        )

    return sheet_names


def _normalize_sheet_name(value: Any) -> str:
    """``1.3``のような数値シート名を文字列へ正規化する。

    シート名は加算の連鎖で求められるため``1.3000000000000003``のような
    浮動小数点誤差を含む。丸めてから文字列化する。
    """
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, (int, float)):
        rounded = round(float(value), 10)
        if rounded.is_integer():
            return str(int(rounded))
        return repr(rounded)
    return str(value).strip()


def find_time_series_columns(workbook_path: str | Path, sheet_name: str) -> dict[str, int]:
    """本せん断シートの Time / Force Y / Moment Y 列と書込み開始行を特定する。

    このシートには見出し行が無いため、せん断応力列の数式が参照している
    Moment Y列をアンカーにして列を決定する(列位置の決め打ちを避ける)。
    """
    path = Path(workbook_path)
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        if sheet_name not in workbook.sheetnames:
            raise ShearAnalysisExcelError(f"シート '{sheet_name}' がありません: {path}")
        worksheet = workbook[sheet_name]

        candidates: list[tuple[int, int, int]] = []
        for row in worksheet.iter_rows(
            min_row=1, max_row=_HEADER_SEARCH_ROWS, max_col=_HEADER_SEARCH_COLUMNS
        ):
            for cell in row:
                formula = cell.value
                if not isinstance(formula, str) or _SHEAR_STRESS_FORMULA_MARKER not in formula:
                    continue
                references = _RELATIVE_REFERENCE_PATTERN.findall(formula)
                referenced_columns = {
                    _column_letter_to_index(letter)
                    for letter, _ in references
                    if f"${letter}$" not in formula
                }
                if len(referenced_columns) != 1:
                    raise ShearAnalysisExcelError(
                        f"シート '{sheet_name}' のせん断応力数式から参照列を一意に特定できません: "
                        f"{formula!r}"
                    )
                candidates.append((cell.row, cell.column, referenced_columns.pop()))
    finally:
        workbook.close()

    if not candidates:
        raise ShearAnalysisExcelError(
            f"シート '{sheet_name}' でせん断応力列(ISBLANKを含む数式)が見つかりません。"
        )

    moment_columns = {moment for _, _, moment in candidates}
    if len(moment_columns) != 1:
        raise ShearAnalysisExcelError(
            f"シート '{sheet_name}' で参照先のMoment Y列が複数見つかりました: "
            f"{sorted(moment_columns)}"
        )

    moment_column = moment_columns.pop()
    start_row = min(row for row, _, _ in candidates)
    formula_column = min(column for _, column, _ in candidates)

    if moment_column < 3:
        raise ShearAnalysisExcelError(
            f"シート '{sheet_name}' のMoment Y列({moment_column})の左に"
            "TimeとForce Yの列が確保できません。"
        )

    columns = {
        "time": moment_column - 2,
        "force_y": moment_column - 1,
        "moment_y": moment_column,
        "shear_stress_formula": formula_column,
        "start_row": start_row,
    }
    if columns["shear_stress_formula"] in {
        columns["time"], columns["force_y"], columns["moment_y"]
    }:
        raise ShearAnalysisExcelError(
            f"シート '{sheet_name}' で書込み対象列とせん断応力数式列が重複しています: {columns}"
        )
    return columns


def _column_letter_to_index(letters: str) -> int:
    index = 0
    for character in letters:
        index = index * 26 + (ord(character) - ord("A") + 1)
    return index


def get_correction_factor(workbook_path: str | Path) -> float | None:
    """「設定・集計」シートの補正係数を取得する。"""
    return _get_labeled_setting(workbook_path, _CORRECTION_FACTOR_LABEL)


def get_peak_ratio(workbook_path: str | Path) -> float | None:
    """「設定・集計」シートの実験ピーク比を取得する。"""
    return _get_labeled_setting(workbook_path, _PEAK_RATIO_LABEL)


def _get_labeled_setting(workbook_path: str | Path, label: str) -> float | None:
    """設定領域のラベルを探し、その右隣セルの数値を返す。"""
    rows = _read_cells(
        Path(workbook_path),
        SETTINGS_SHEET_NAME,
        _HEADER_SEARCH_ROWS + _MAX_CONDITION_ROWS,
    )
    for row_index in range(min(_HEADER_SEARCH_ROWS, len(rows))):
        for column_index, value in enumerate(rows[row_index]):
            if isinstance(value, str) and value.strip() == label:
                setting = _cell(rows, row_index, column_index + 1)
                return float(setting) if isinstance(setting, (int, float)) else None
    return None


def get_steady_slope_cells(workbook_path: str | Path, condition_id: int) -> dict[str, str]:
    """条件番号に対応する定常傾き(補正前・補正後)のセル番地を返す。"""
    path = Path(workbook_path)
    rows = _read_cells(path, SETTINGS_SHEET_NAME, _HEADER_SEARCH_ROWS + _MAX_CONDITION_ROWS)
    _, condition_row = _locate_condition_row(rows, condition_id)

    slope_column = None
    for row_index in range(min(_HEADER_SEARCH_ROWS, len(rows))):
        slope_column = _find_header_column(rows, row_index, _STEADY_SLOPE_HEADER)
        if slope_column is not None:
            slope_header_row = row_index
            break
    if slope_column is None:
        raise ShearAnalysisExcelError(
            f"「{SETTINGS_SHEET_NAME}」シートに見出し '{_STEADY_SLOPE_HEADER}' が見つかりません。"
        )

    raw_column = _find_header_column(rows, slope_header_row + 1, _STEADY_SLOPE_RAW_HEADER)
    corrected_column = _find_header_column(
        rows, slope_header_row + 1, _STEADY_SLOPE_CORRECTED_HEADER
    )
    if raw_column is None or corrected_column is None:
        raise ShearAnalysisExcelError(
            f"「{SETTINGS_SHEET_NAME}」シートに '{_STEADY_SLOPE_RAW_HEADER}' / "
            f"'{_STEADY_SLOPE_CORRECTED_HEADER}' の見出しが見つかりません。"
        )

    excel_row = condition_row + 1
    return {
        "raw": f"{_index_to_column_letter(raw_column + 1)}{excel_row}",
        "corrected": f"{_index_to_column_letter(corrected_column + 1)}{excel_row}",
    }


def _index_to_column_letter(index: int) -> str:
    letters = ""
    while index > 0:
        index, remainder = divmod(index - 1, 26)
        letters = chr(ord("A") + remainder) + letters
    return letters


def update_and_recalculate(
    workbook_path: str | Path,
    condition_id: int,
    time_series_by_load: dict[float, tuple[list[float], list[float], list[float]]],
) -> dict[str, Any]:
    """本せん断の時系列を書き込み、ブックを再計算して結果を取得する。

    グラフ・書式・配列数式を保持するためExcel COMで操作する。COMが使えない場合は
    書込みを行わず``recalculated=False``を返し、古いキャッシュ値を結果としない。
    """
    path = Path(workbook_path).resolve()
    if not path.is_file():
        raise FileNotFoundError(f"解析Excelが見つかりません: {path}")

    sheet_names = get_shear_sheet_names(path, condition_id)
    columns_by_load = {
        load_kpa: find_time_series_columns(path, sheet_names[load_kpa])
        for load_kpa in time_series_by_load
    }
    slope_cells = get_steady_slope_cells(path, condition_id)

    result: dict[str, Any] = {
        "workbook": str(path),
        "condition_id": condition_id,
        "sheet_names": {str(load): name for load, name in sheet_names.items()},
        "columns": {str(load): columns for load, columns in columns_by_load.items()},
        "steady_slope_cells": slope_cells,
        "correction_factor": get_correction_factor(path),
        "recalculated": False,
    }

    if not is_recalculation_available():
        result["recalculation_required"] = True
        result["note"] = (
            "Excel COM(pywin32)が利用できないため書込みと再計算を行っていません。"
            "古いキャッシュ値は結果として採用しません。"
        )
        return result

    import pythoncom
    import win32com.client

    pythoncom.CoInitialize()
    excel = None
    workbook = None
    try:
        excel = win32com.client.DispatchEx("Excel.Application")
        excel.Visible = False
        excel.DisplayAlerts = False
        workbook = excel.Workbooks.Open(str(path))

        written: dict[str, int] = {}
        for load_kpa, (times, forces, moments) in time_series_by_load.items():
            sheet_name = sheet_names[load_kpa]
            columns = columns_by_load[load_kpa]
            worksheet = workbook.Worksheets(sheet_name)
            _write_time_series_com(worksheet, columns, times, forces, moments)
            written[str(load_kpa)] = len(times)

        excel.Application.CalculateFullRebuild()
        workbook.Save()

        result["written_rows"] = written
        result["recalculated"] = True

        settings_sheet = workbook.Worksheets(SETTINGS_SHEET_NAME)
        slope_raw, slope_raw_error = _cell_value_com(settings_sheet, slope_cells["raw"])
        slope_corrected, slope_corrected_error = _cell_value_com(
            settings_sheet, slope_cells["corrected"]
        )
        result["steady_slope_raw"] = slope_raw
        result["steady_slope_corrected"] = slope_corrected
        steady_shear_stress = _read_steady_shear_stress_com(workbook, slope_cells["raw"])
        steady_errors = steady_shear_stress.pop("__errors__", {})
        result["steady_shear_stress_kpa"] = steady_shear_stress

        errors = {
            key: error
            for key, error in (
                ("steady_slope_raw", slope_raw_error),
                ("steady_slope_corrected", slope_corrected_error),
            )
            if error is not None
        }
        errors.update(
            {
                f"steady_shear_stress_kpa[{load}]": error
                for load, error in steady_errors.items()
            }
        )
        if errors:
            result["formula_errors"] = errors
            result["recalculation_required"] = True
            result["note"] = (
                "再計算後も数式エラーが残っています。未実行の荷重がある場合は"
                "3/5/7 kPaすべてを実行してください。"
            )
    finally:
        if workbook is not None:
            workbook.Close(SaveChanges=False)
        if excel is not None:
            excel.Quit()
        pythoncom.CoUninitialize()

    return result


def _write_time_series_com(
    worksheet: Any,
    columns: dict[str, int],
    times: list[float],
    forces: list[float],
    moments: list[float],
) -> None:
    """Time / Force Y / Moment Y を書き込む。せん断応力の数式列には触れない。"""
    start_row = columns["start_row"]
    used_rows = max(worksheet.UsedRange.Row + worksheet.UsedRange.Rows.Count - 1, start_row)

    for key in ("time", "force_y", "moment_y"):
        column = columns[key]
        worksheet.Range(
            worksheet.Cells(start_row, column), worksheet.Cells(used_rows, column)
        ).ClearContents()

    end_row = start_row + len(times) - 1
    worksheet.Range(
        worksheet.Cells(start_row, columns["time"]), worksheet.Cells(end_row, columns["time"])
    ).Value = [[value] for value in times]
    worksheet.Range(
        worksheet.Cells(start_row, columns["force_y"]),
        worksheet.Cells(end_row, columns["force_y"]),
    ).Value = [[value] for value in forces]
    worksheet.Range(
        worksheet.Cells(start_row, columns["moment_y"]),
        worksheet.Cells(end_row, columns["moment_y"]),
    ).Value = [[value] for value in moments]


def _cell_value_com(worksheet: Any, address: str) -> tuple[float | None, str | None]:
    """セル値を返す。数式エラーの場合は値をNoneとし、エラー名を併せて返す。"""
    return _coerce_excel_value(worksheet.Range(address).Value)


def _coerce_excel_value(value: Any) -> tuple[float | None, str | None]:
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        return None, None
    error_name = _EXCEL_ERROR_VALUES.get(int(value))
    if error_name is not None:
        return None, error_name
    return float(value), None


def _read_steady_shear_stress_com(workbook: Any, slope_cell: str) -> dict[str, float | None]:
    """定常傾きセルの参照先(グラフシート)から各荷重の定常せん断応力を読む。"""
    formula = workbook.Worksheets(SETTINGS_SHEET_NAME).Range(slope_cell).Formula
    match = re.search(r"=\s*'?([^'!]+)'?!\$?([A-Z]{1,3})\$?(\d+)", str(formula))
    if match is None:
        raise ShearAnalysisExcelError(
            f"定常傾きセル {slope_cell} の参照先を解析できません: {formula!r}"
        )

    graph_sheet_name, slope_column_letter, slope_row_text = match.groups()
    graph_sheet = workbook.Worksheets(graph_sheet_name)
    slope_column = _column_letter_to_index(slope_column_letter)
    slope_row = int(slope_row_text)

    # 定常せん断応力は定常傾きの1行上、荷重は同ブロックの数行下に並ぶ。
    steady_row = slope_row - 1
    loads_row = _find_loads_row_com(graph_sheet, slope_column, slope_row)

    steady: dict[str, float | None] = {}
    errors: dict[str, str] = {}
    for offset in range(len(_SHEAR_SHEET_HEADERS)):
        column = slope_column + offset
        load_value = graph_sheet.Cells(loads_row, column).Value
        value, error_name = _coerce_excel_value(graph_sheet.Cells(steady_row, column).Value)
        key = f"{float(load_value):g}" if isinstance(load_value, (int, float)) else str(column)
        steady[key] = value
        if error_name is not None:
            errors[key] = error_name
    if errors:
        steady["__errors__"] = errors  # type: ignore[assignment]
    return steady


def _find_loads_row_com(graph_sheet: Any, slope_column: int, slope_row: int) -> int:
    """荷重(3, 5, 7)が並ぶ行を探し、列対応の妥当性を確認する。"""
    expected = sorted(_SHEAR_SHEET_HEADERS)
    for row in range(slope_row + 1, slope_row + 12):
        values = [
            graph_sheet.Cells(row, slope_column + offset).Value
            for offset in range(len(expected))
        ]
        if all(isinstance(value, (int, float)) for value in values) and [
            float(value) for value in values
        ] == expected:
            return row
    raise ShearAnalysisExcelError(
        f"グラフシートで荷重{expected}が並ぶ行を特定できませんでした"
        f"(定常傾き列={_index_to_column_letter(slope_column)}, 行={slope_row})。"
    )
