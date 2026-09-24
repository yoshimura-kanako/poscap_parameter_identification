"""せん断試験条件表作成のユニットテスト。"""

from pathlib import Path

from openpyxl import load_workbook

from poscap_calibration.workflows.phase1_shear import create_shear_analysis_workbook


TEMPLATE = Path(__file__).parents[2] / "templates" / "excel" / "shear_analysis_template.xlsx"


def test_create_shear_analysis_workbook_copies_and_sets_peak_ratio(tmp_path: Path) -> None:
    result_path = create_shear_analysis_workbook(
        "TEST_PROJECT",
        1.25,
        project_root=tmp_path / "projects",
        template_path=TEMPLATE,
    )

    work_path = tmp_path / "projects" / "TEST_PROJECT" / "shear_test" / "work" / "shear_analysis_working.xlsx"
    log_path = tmp_path / "projects" / "TEST_PROJECT" / "shear_test" / "step1_shear_analysis.log"

    assert work_path.is_file()
    assert result_path == tmp_path / "projects" / "TEST_PROJECT" / "shear_test" / "shear_analysis.xlsx"
    assert result_path.is_file()
    assert log_path.is_file()
    assert "peak_ratio=1.25" in log_path.read_text(encoding="utf-8")

    for path in (work_path, result_path):
        workbook = load_workbook(path, read_only=True, data_only=False)
        assert workbook["設定・集計"]["B3"].value == 1.25
        workbook.close()

    original = load_workbook(TEMPLATE, read_only=True, data_only=False)
    assert original["設定・集計"]["B3"].value is None
    original.close()