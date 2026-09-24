"""壁面摩擦試験の条件表作成(手順書Phase2 1.2.4)のユニットテスト。"""

from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from poscap_calibration.wall_friction_analysis_excel import (
    WallFrictionExcelError,
    is_recalculation_available,
    read_wall_friction_conditions,
)
from poscap_calibration.workflows.phase2_wall import (
    create_wall_friction_analysis_workbook,
)

TEMPLATE = (
    Path(__file__).parents[2]
    / "templates"
    / "excel"
    / "wall_friction_analysis_template.xlsx"
)

requires_excel = pytest.mark.skipif(
    not is_recalculation_available(), reason="Excel COM(pywin32)が利用できません。"
)


def _create_settings_workbook(path: Path, rows: list[tuple[float, float, float]]) -> Path:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "設定・集計"
    worksheet["D1"] = "条件"
    worksheet["E1"] = "壁面パラメータ"
    worksheet["E2"] = "動摩擦係数"
    worksheet["F2"] = "静止摩擦係数"
    for offset, (condition_id, dynamic, static) in enumerate(rows):
        row = 3 + offset
        worksheet.cell(row=row, column=4, value=condition_id)
        worksheet.cell(row=row, column=5, value=dynamic)
        worksheet.cell(row=row, column=6, value=static)
    workbook.save(path)
    workbook.close()
    return path


def test_read_wall_friction_conditions(tmp_path: Path) -> None:
    path = _create_settings_workbook(
        tmp_path / "wall.xlsx",
        [(1, 0.11291, 0.11291), (2, 0.22582, 0.22582), (3, 0.33873, 0.33873)],
    )

    conditions = read_wall_friction_conditions(path)

    assert [c.condition_id for c in conditions] == [1, 2, 3]
    assert conditions[1].dynamic_friction == pytest.approx(0.22582)
    assert conditions[2].static_friction == pytest.approx(0.33873)


def test_read_wall_friction_conditions_rejects_formula_errors(tmp_path: Path) -> None:
    path = _create_settings_workbook(
        tmp_path / "wall.xlsx", [(1, 0.11291, 0.11291)]
    )
    workbook = load_workbook(path)
    workbook["設定・集計"]["E3"] = "#DIV/0!"
    workbook.save(path)
    workbook.close()

    with pytest.raises(WallFrictionExcelError):
        read_wall_friction_conditions(path)


def test_read_wall_friction_conditions_requires_settings_sheet(tmp_path: Path) -> None:
    workbook = Workbook()
    workbook.active.title = "別のシート"
    path = tmp_path / "wall.xlsx"
    workbook.save(path)
    workbook.close()

    with pytest.raises(WallFrictionExcelError):
        read_wall_friction_conditions(path)


def test_create_wall_friction_analysis_workbook_rejects_invalid_inputs(
    tmp_path: Path,
) -> None:
    for peak_ratio, steady_slope in ((0.0, 0.22582), (1.0813, -1.0)):
        with pytest.raises(ValueError):
            create_wall_friction_analysis_workbook(
                "TEST_PROJECT",
                peak_ratio,
                steady_slope,
                project_root=tmp_path / "projects",
                template_path=TEMPLATE,
            )


@requires_excel
def test_create_wall_friction_analysis_workbook(tmp_path: Path) -> None:
    workbook_path, conditions = create_wall_friction_analysis_workbook(
        "TEST_PROJECT",
        1.0813,
        0.22582,
        project_root=tmp_path / "projects",
        template_path=TEMPLATE,
    )

    test_directory = tmp_path / "projects" / "TEST_PROJECT" / "wall_friction_test"
    assert workbook_path == test_directory / "wall_friction_analysis.xlsx"
    assert workbook_path.is_file()

    log_text = (test_directory / "step1_wall_friction_analysis.log").read_text(
        encoding="utf-8"
    )
    assert "peak_ratio=1.0813" in log_text
    assert "experiment_steady_slope=0.22582" in log_text
    assert "condition=2 wall_dynamic_friction=0.225820" in log_text

    assert [c.condition_id for c in conditions] == [1, 2, 3]
    assert conditions[0].dynamic_friction == pytest.approx(0.22582 * 0.5)
    assert conditions[1].dynamic_friction == pytest.approx(0.22582)
    assert conditions[2].dynamic_friction == pytest.approx(0.22582 * 1.5)
    for condition in conditions:
        assert condition.static_friction == pytest.approx(condition.dynamic_friction)

    # テンプレート原本は変更しない。
    original = load_workbook(TEMPLATE, read_only=True, data_only=False)
    assert original["設定・集計"]["B3"].value is None
    assert original["設定・集計"]["B5"].value is None
    original.close()
