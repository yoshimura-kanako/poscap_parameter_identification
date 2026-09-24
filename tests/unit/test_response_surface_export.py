"""応答曲面用データ出力(手順書1.2.8 / 2.2.7)のユニットテスト。"""

from __future__ import annotations

import csv
from pathlib import Path

import pytest
from openpyxl import Workbook

from poscap_calibration.response_surface_export import (
    RESPONSE_SURFACE_SHEET_NAME,
    ResponseSurfaceExportError,
    export_response_surface_data,
)


def _create_workbook(path: Path, rows: list[list[object]]) -> Path:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = RESPONSE_SURFACE_SHEET_NAME
    for row in rows:
        worksheet.append(row)
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    workbook.close()
    return path


def _condition_rows(count: int) -> list[list[object]]:
    header: list[object] = ["転がり抵抗", "動摩擦係数", "定常傾き"]
    rows = [header]
    for index in range(count):
        rows.append([0.2 + index * 0.05, 0.2 + index * 0.05, 1.5 + index * 0.1])
    return rows


def test_export_response_surface_data_writes_whole_sheet(tmp_path: Path) -> None:
    workbook_path = _create_workbook(tmp_path / "shear_analysis.xlsx", _condition_rows(9))
    output_csv = tmp_path / "results" / "shear_response_surface_data.csv"

    result = export_response_surface_data(
        workbook_path, output_csv, recalculate=False
    )

    assert output_csv.is_file()
    assert result["row_count"] == 10
    assert result["data_row_count"] == 9
    assert result["column_count"] == 3
    assert result["sheet"] == RESPONSE_SURFACE_SHEET_NAME

    with output_csv.open(encoding="utf-8-sig", newline="") as csv_file:
        rows = list(csv.reader(csv_file))
    assert rows[0] == ["転がり抵抗", "動摩擦係数", "定常傾き"]
    assert len(rows) == 10
    assert rows[1][2] == "1.5"


def test_export_response_surface_data_handles_more_than_nine_conditions(
    tmp_path: Path,
) -> None:
    """条件数が27に増えても行数を固定せずシート全体を出力する。"""
    workbook_path = _create_workbook(tmp_path / "shear_analysis.xlsx", _condition_rows(27))
    output_csv = tmp_path / "response_surface.csv"

    result = export_response_surface_data(workbook_path, output_csv, recalculate=False)

    assert result["data_row_count"] == 27
    with output_csv.open(encoding="utf-8-sig", newline="") as csv_file:
        rows = list(csv.reader(csv_file))
    assert len(rows) == 28


def test_export_response_surface_data_rejects_formula_errors(tmp_path: Path) -> None:
    rows = _condition_rows(2)
    rows[2][2] = "#VALUE!"
    workbook_path = _create_workbook(tmp_path / "shear_analysis.xlsx", rows)
    output_csv = tmp_path / "response_surface.csv"

    with pytest.raises(ResponseSurfaceExportError):
        export_response_surface_data(workbook_path, output_csv, recalculate=False)

    assert not output_csv.exists()


def test_export_response_surface_data_allows_incomplete_when_requested(
    tmp_path: Path,
) -> None:
    rows = _condition_rows(2)
    rows[2][2] = "#VALUE!"
    workbook_path = _create_workbook(tmp_path / "shear_analysis.xlsx", rows)
    output_csv = tmp_path / "response_surface.csv"

    result = export_response_surface_data(
        workbook_path, output_csv, recalculate=False, allow_incomplete=True
    )

    assert output_csv.is_file()
    assert result["formula_errors"] == ["R3C3=#VALUE!"]


def test_export_response_surface_data_raises_when_sheet_missing(tmp_path: Path) -> None:
    workbook = Workbook()
    workbook.active.title = "別のシート"
    workbook_path = tmp_path / "shear_analysis.xlsx"
    workbook.save(workbook_path)
    workbook.close()

    with pytest.raises(ResponseSurfaceExportError):
        export_response_surface_data(
            workbook_path, tmp_path / "out.csv", recalculate=False
        )


def _saor_condition_rows(count: int) -> list[list[object]]:
    header: list[object] = ["転がり抵抗", "動摩擦係数", "安息角"]
    rows = [header]
    for index in range(count):
        rows.append([0.2 + index * 0.05, 0.2 + index * 0.05, 33.2 + index])
    return rows


def test_export_saor_response_surface_data(tmp_path: Path) -> None:
    """安息角試験(手順書2.2.7)でも同じ処理を流用できることを確認する。"""
    workbook_path = _create_workbook(
        tmp_path / "saor_analysis.xlsx", _saor_condition_rows(9)
    )
    output_csv = tmp_path / "saor_response_surface_data.csv"

    result = export_response_surface_data(workbook_path, output_csv, recalculate=False)

    assert result["row_count"] == 10
    assert result["data_row_count"] == 9
    assert result["column_count"] == 3

    with output_csv.open(encoding="utf-8-sig", newline="") as csv_file:
        rows = list(csv.reader(csv_file))
    assert rows[0] == ["転がり抵抗", "動摩擦係数", "安息角"]
    assert len(rows) == 10


def test_export_response_surface_data_accepts_custom_sheet_name(tmp_path: Path) -> None:
    workbook = Workbook()
    workbook.active.title = "別のシート"
    for row in _saor_condition_rows(2):
        workbook.active.append(row)
    workbook_path = tmp_path / "saor_analysis.xlsx"
    workbook.save(workbook_path)
    workbook.close()

    result = export_response_surface_data(
        workbook_path,
        tmp_path / "out.csv",
        recalculate=False,
        sheet_name="別のシート",
    )

    assert result["sheet"] == "別のシート"
    assert result["data_row_count"] == 2
