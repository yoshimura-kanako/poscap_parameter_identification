"""安息角実験データ点のExcel転記のユニットテスト。"""

from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from poscap_calibration.saor_analysis_excel import (
    EXPERIMENT_POINT_COLUMNS,
    SaorAnalysisExcelError,
    get_condition_sheet_name,
    load_experiment_points,
    write_experiment_points,
)

WORKBOOK = (
    Path(__file__).parents[2] / "projects" / "0" / "saor_test" / "saor_analysis.xlsx"
)


def _create_workbook(tmp_path: Path, condition_count: int = 9) -> Path:
    workbook = Workbook()
    settings = workbook.active
    settings.title = "設定・集計"
    settings["D1"] = "条件"
    settings["E1"] = "転がり抵抗"
    settings["F1"] = "動摩擦係数"
    settings["G1"] = "静止摩擦係数"
    settings["H1"] = "シート名"

    for index in range(condition_count):
        row = 2 + index
        settings.cell(row=row, column=4, value=index + 1)
        settings.cell(row=row, column=8, value=index + 1)
        sheet = workbook.create_sheet(str(index + 1))
        # 転記対象外の列に数式があっても壊さないことを確認するため配置する。
        sheet["Q2"] = "=SLOPE(N2:N100,K2:K100)"
        sheet["P3"] = "安息角"

    path = tmp_path / "saor_analysis.xlsx"
    workbook.save(path)
    workbook.close()
    return path


def _create_csv(tmp_path: Path, rows: int = 3) -> Path:
    csv_path = tmp_path / "experiment_data_points.csv"
    lines = [",".join(EXPERIMENT_POINT_COLUMNS)]
    for index in range(rows):
        lines.append(",".join(str(index + offset * 0.1) for offset in range(6)))
    csv_path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return csv_path


def test_get_condition_sheet_name_uses_settings_sheet(tmp_path: Path) -> None:
    workbook_path = _create_workbook(tmp_path)

    assert get_condition_sheet_name(workbook_path, 1) == "1"
    assert get_condition_sheet_name(workbook_path, 9) == "9"


def test_get_condition_sheet_name_raises_for_unknown_condition(tmp_path: Path) -> None:
    workbook_path = _create_workbook(tmp_path, condition_count=2)

    with pytest.raises(SaorAnalysisExcelError):
        get_condition_sheet_name(workbook_path, 5)


def test_load_experiment_points_raises_when_column_missing(tmp_path: Path) -> None:
    csv_path = tmp_path / "bad.csv"
    csv_path.write_text("x_coord,min_y_coordinate\n0.0,1.0\n", encoding="utf-8")

    with pytest.raises(SaorAnalysisExcelError) as error:
        load_experiment_points(csv_path, 3)

    message = str(error.value)
    assert "条件3" in message
    assert "max_y_coordinate" in message


def test_load_experiment_points_raises_when_file_missing(tmp_path: Path) -> None:
    with pytest.raises(FileNotFoundError):
        load_experiment_points(tmp_path / "missing.csv", 1)


def test_write_experiment_points_writes_header_and_values(tmp_path: Path) -> None:
    workbook_path = _create_workbook(tmp_path)
    csv_path = _create_csv(tmp_path, rows=3)

    result = write_experiment_points(workbook_path, 1, csv_path)

    assert result["sheet_name"] == "1"
    assert result["row_count"] == 3

    workbook = load_workbook(workbook_path)
    sheet = workbook["1"]
    assert [sheet.cell(row=1, column=i + 1).value for i in range(6)] == list(
        EXPERIMENT_POINT_COLUMNS
    )
    assert sheet["A2"].value == pytest.approx(0.0)
    assert sheet["F2"].value == pytest.approx(0.5)
    assert sheet["A4"].value == pytest.approx(2.0)
    # 転記対象外の数式は保持される。
    assert sheet["Q2"].value == "=SLOPE(N2:N100,K2:K100)"
    workbook.close()


def test_write_experiment_points_clears_old_rows(tmp_path: Path) -> None:
    workbook_path = _create_workbook(tmp_path)
    write_experiment_points(workbook_path, 1, _create_csv(tmp_path, rows=5))

    shorter_csv = tmp_path / "shorter.csv"
    shorter_csv.write_text(
        ",".join(EXPERIMENT_POINT_COLUMNS) + "\n" + ",".join(["1.0"] * 6) + "\n",
        encoding="utf-8",
    )
    write_experiment_points(workbook_path, 1, shorter_csv)

    workbook = load_workbook(workbook_path)
    sheet = workbook["1"]
    assert sheet["A2"].value == pytest.approx(1.0)
    assert sheet["A3"].value is None
    workbook.close()


def test_write_experiment_points_targets_condition_specific_sheet(tmp_path: Path) -> None:
    workbook_path = _create_workbook(tmp_path)
    csv_path = _create_csv(tmp_path)

    write_experiment_points(workbook_path, 7, csv_path)

    workbook = load_workbook(workbook_path)
    assert workbook["7"]["A1"].value == "x_coord"
    assert workbook["1"]["A1"].value is None
    workbook.close()


def test_write_experiment_points_supports_all_nine_conditions(tmp_path: Path) -> None:
    workbook_path = _create_workbook(tmp_path)
    csv_path = _create_csv(tmp_path)

    for condition_id in range(1, 10):
        result = write_experiment_points(workbook_path, condition_id, csv_path)
        assert result["sheet_name"] == str(condition_id)

    workbook = load_workbook(workbook_path)
    for condition_id in range(1, 10):
        assert workbook[str(condition_id)]["A1"].value == "x_coord"
    workbook.close()


@pytest.mark.skipif(not WORKBOOK.is_file(), reason="解析Excelが存在しない環境ではスキップ")
def test_real_workbook_resolves_sheet_names() -> None:
    assert get_condition_sheet_name(WORKBOOK, 1) == "1"
    assert get_condition_sheet_name(WORKBOOK, 9) == "9"
