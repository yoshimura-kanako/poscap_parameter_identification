"""9条件バッチ実行のユニットテスト(MockRockyClient使用、実シミュレーションなし)。"""

from __future__ import annotations

import json
import shutil
from pathlib import Path

import pytest
from openpyxl import Workbook

from poscap_calibration.rocky.mock_client import MockRockyClient
from poscap_calibration.workflows.shear_batch import run_condition, run_conditions

RAW_CSV_HEADER = "Particle ID,Coordinate : X,Coordinate : Y,Coordinate : Z,Particle Size"
LOADS = (3.0, 5.0, 7.0)
GEOMETRY = "FT4_ShearCell2_5deg_mm"


@pytest.fixture(name="particles_csv")
def fixture_particles_csv(tmp_path_factory: pytest.TempPathFactory) -> Path:
    """MockRockyClientが返す粒子CSV(小規模、変換処理の検証には十分)。"""
    path = tmp_path_factory.mktemp("particles") / "particles_raw.csv"
    rows = "\n".join(
        f"{index},0.00{index}1,0.00{index}2,0.00{index}3,0.0001" for index in range(1, 6)
    )
    path.write_text(f"{RAW_CSV_HEADER}\n{rows}\n", encoding="utf-8")
    return path


def _create_project(path: Path, content: str = "template") -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(content, encoding="utf-8")
    files_dir = path.with_suffix(f"{path.suffix}.files")
    files_dir.mkdir(exist_ok=True)
    (files_dir / "data.bin").write_text("data", encoding="utf-8")
    return path


def _create_workbook(path: Path, condition_count: int = 9) -> Path:
    """「設定・集計」シートを持つ条件表を作成する(9条件)。"""
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "設定・集計"
    worksheet["D1"] = "条件"
    worksheet["E2"] = "転がり抵抗"
    worksheet["F2"] = "動摩擦係数"
    worksheet["G2"] = "静止摩擦係数"

    rolling_levels = [0.2, 0.45, 0.7]
    dynamic_levels = [0.2, 0.45, 0.7]
    for index in range(condition_count):
        row = 3 + index
        worksheet.cell(row=row, column=4, value=index + 1)
        worksheet.cell(row=row, column=5, value=rolling_levels[index // 3])
        worksheet.cell(row=row, column=6, value=dynamic_levels[index % 3])
        worksheet.cell(row=row, column=7, value=round(dynamic_levels[index % 3] * 1.0983, 5))

    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    workbook.close()
    return path


def _build_environment(tmp_path: Path) -> dict[str, object]:
    filling_base = _create_project(tmp_path / "filling" / "Filling.rocky")
    pre_shear_template = _create_project(tmp_path / "templates" / "Pre_shear.rocky")
    shear_templates = {
        load: _create_project(tmp_path / "templates" / f"Shear_{load:g}kPa.rocky")
        for load in LOADS
    }
    workbook_path = _create_workbook(tmp_path / "results" / "shear_analysis.xlsx")
    return {
        "filling_base_project_path": filling_base,
        "pre_shear_template_path": pre_shear_template,
        "shear_templates": shear_templates,
        "workbook_path": workbook_path,
        "output_root": tmp_path / "results",
        "shear_cell_geometry_name": GEOMETRY,
    }


def _client_factory_for(particles_csv: Path):
    def factory() -> MockRockyClient:
        return MockRockyClient(completed=True, particles_csv=particles_csv)

    return factory


def test_run_conditions_processes_all_nine_conditions(
    tmp_path: Path, particles_csv: Path
) -> None:
    environment = _build_environment(tmp_path)

    batch = run_conditions(
        _client_factory_for(particles_csv),
        loads_kpa=LOADS,
        update_excel=False,
        **environment,
    )

    assert batch["condition_ids"] == list(range(1, 10))
    assert batch["completed_condition_ids"] == list(range(1, 10))
    assert batch["failed_condition_ids"] == []

    output_root = Path(environment["output_root"])
    assert (output_root / "shear_batch_summary.json").is_file()

    for condition_id in range(1, 10):
        condition_directory = output_root / f"condition_{condition_id:02d}"
        assert (condition_directory / "fill" / "status.json").is_file()
        assert (condition_directory / "pre_shear" / "status.json").is_file()
        assert (condition_directory / "shear_summary.json").is_file()
        for load in LOADS:
            csv_path = (
                condition_directory
                / f"shear_{load:g}kpa"
                / "raw"
                / f"{load:g}kPa_ShearTest_time.csv"
            )
            assert csv_path.is_file()


def test_run_conditions_uses_condition_specific_parameters(
    tmp_path: Path, particles_csv: Path
) -> None:
    environment = _build_environment(tmp_path)

    run_conditions(
        _client_factory_for(particles_csv),
        condition_ids=(1, 9),
        loads_kpa=(3.0,),
        update_excel=False,
        **environment,
    )

    output_root = Path(environment["output_root"])
    first = json.loads(
        (output_root / "condition_01" / "shear_summary.json").read_text(encoding="utf-8")
    )
    ninth = json.loads(
        (output_root / "condition_09" / "shear_summary.json").read_text(encoding="utf-8")
    )

    assert first["parameters"]["rolling_resistance"] == pytest.approx(0.2)
    assert ninth["parameters"]["rolling_resistance"] == pytest.approx(0.7)
    assert ninth["parameters"]["dynamic_friction"] == pytest.approx(0.7)
    assert not (output_root / "condition_02").exists()


def test_run_conditions_skips_completed_steps(tmp_path: Path, particles_csv: Path) -> None:
    environment = _build_environment(tmp_path)
    factory = _client_factory_for(particles_csv)

    run_conditions(
        factory,
        condition_ids=(1,),
        loads_kpa=(3.0,),
        update_excel=False,
        **environment,
    )

    second_batch = run_conditions(
        factory,
        condition_ids=(1,),
        loads_kpa=(3.0,),
        update_excel=False,
        **environment,
    )

    steps = json.loads(
        (Path(environment["output_root"]) / "condition_01" / "shear_summary.json").read_text(
            encoding="utf-8"
        )
    )["steps"]
    assert steps["fill"]["status"] == "skipped"
    assert steps["pre_shear"]["status"] == "skipped"
    assert steps["shear_3kpa"]["status"] == "skipped"
    assert second_batch["completed_condition_ids"] == [1]


def test_run_conditions_isolates_failures(tmp_path: Path, particles_csv: Path) -> None:
    environment = _build_environment(tmp_path)
    factory = _client_factory_for(particles_csv)
    output_root = Path(environment["output_root"])

    batch = run_conditions(
        factory,
        condition_ids=(1, 2, 3),
        loads_kpa=(3.0,),
        update_excel=False,
        **environment,
    )
    assert batch["completed_condition_ids"] == [1, 2, 3]

    # 条件2だけ充填成果物を欠損させ、その条件のみ失敗することを確認する。
    shutil.rmtree(output_root / "condition_02")
    (output_root / "condition_02" / "fill").mkdir(parents=True, exist_ok=True)
    (output_root / "condition_02" / "fill" / "status.json").write_text(
        json.dumps({"status": "completed"}), encoding="utf-8"
    )

    batch = run_conditions(
        factory,
        condition_ids=(1, 2, 3),
        loads_kpa=(3.0,),
        update_excel=False,
        **environment,
    )
    assert batch["failed_condition_ids"] == [2]
    assert batch["completed_condition_ids"] == [1, 3]


def test_run_condition_records_failure_in_summary(
    tmp_path: Path, particles_csv: Path
) -> None:
    environment = _build_environment(tmp_path)
    environment["filling_base_project_path"] = tmp_path / "missing.rocky"

    from poscap_calibration.condition_table import get_shear_condition

    condition = get_shear_condition(environment["workbook_path"], 1)

    with pytest.raises(FileNotFoundError):
        run_condition(
            _client_factory_for(particles_csv),
            condition,
            loads_kpa=(3.0,),
            update_excel=False,
            **environment,
        )

    summary = json.loads(
        (Path(environment["output_root"]) / "condition_01" / "shear_summary.json").read_text(
            encoding="utf-8"
        )
    )
    assert summary["status"] == "failed"
    assert "error" in summary
